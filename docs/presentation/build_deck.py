"""Generate docs/presentation/PeerDesk_Architecture.pptx (5 slides).

Run: python3 docs/presentation/build_deck.py
Needs: pip install python-pptx
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

NAVY = RGBColor(0x1F, 0x2A, 0x44)
INK = RGBColor(0x22, 0x28, 0x33)
MUTED = RGBColor(0x5F, 0x6B, 0x7A)
LIGHT = RGBColor(0xF3, 0xF5, 0xF8)
LINE = RGBColor(0x9A, 0xA4, 0xB1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLUE = RGBColor(0x2F, 0x6F, 0xDE)   # viewer / client
BLUE_BG = RGBColor(0xE8, 0xF0, 0xFD)
ORANGE = RGBColor(0xD9, 0x6C, 0x1A)  # host / server
ORANGE_BG = RGBColor(0xFD, 0xF0, 0xE6)
PURPLE = RGBColor(0x6E, 0x4F, 0xC0)  # shared lib
PURPLE_BG = RGBColor(0xF1, 0xED, 0xFB)
TEAL = RGBColor(0x0E, 0x8F, 0x7E)    # network / security
TEAL_BG = RGBColor(0xE6, 0xF6, 0xF3)
RED = RGBColor(0xC2, 0x3B, 0x3B)

FONT = "Calibri"
OUT = Path(__file__).with_name("PeerDesk_Architecture.pptx")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


# ---------- helpers ----------

def text(slide, x, y, w, h, runs, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font=FONT):
    """runs: str, or list of paragraphs; a paragraph is str or list of (text, {opts})."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    _fill_tf(tf, runs, size, color, bold, align, font)
    return tb


def _fill_tf(tf, runs, size, color, bold, align, font=FONT):
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        segs = para if isinstance(para, list) else [(para, {})]
        for seg, opts in segs:
            r = p.add_run()
            r.text = seg
            r.font.name = opts.get("font", font)
            r.font.size = Pt(opts.get("size", size))
            r.font.bold = opts.get("bold", bold)
            r.font.italic = opts.get("italic", False)
            r.font.color.rgb = opts.get("color", color)
        if "space" in (segs[0][1] if segs else {}):
            p.space_before = Pt(segs[0][1]["space"])


def box(slide, x, y, w, h, fill=WHITE, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        runs=None, size=12, color=INK, bold=False, align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE, dash=False, line_w=1.0, radius=0.08):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
        if dash:
            s.line.dash_style = 7  # dash
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    if runs is not None:
        _fill_tf(tf, runs, size, color, bold, align)
    return s


def line(slide, x1, y1, x2, y2, color=LINE, width=1.25, head=None, tail=None, dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                   Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = 7
    ln = c.line._get_or_add_ln()
    for tag, kind in (("a:headEnd", head), ("a:tailEnd", tail)):
        if kind:
            el = etree.SubElement(ln, qn(tag))
            el.set("type", kind)
            el.set("w", "med")
            el.set("len", "med")
    return c


def arrow(slide, x1, y1, x2, y2, color=INK, width=1.5, dash=False, both=False):
    return line(slide, x1, y1, x2, y2, color=color, width=width, tail="triangle",
                head="triangle" if both else None, dash=dash)


def header(slide, n, title, subtitle):
    bar = box(slide, 0, 0, 13.333, 1.05, fill=NAVY, line=None, shape=MSO_SHAPE.RECTANGLE)
    bar.text_frame.text = ""
    text(slide, 0.5, 0.14, 11, 0.55, title, size=28, color=WHITE, bold=True)
    text(slide, 0.5, 0.62, 11.5, 0.35, subtitle, size=14, color=RGBColor(0xC9, 0xD3, 0xE3))
    text(slide, 12.2, 0.3, 0.8, 0.45, f"{n} / 5", size=13, color=RGBColor(0xC9, 0xD3, 0xE3),
         align=PP_ALIGN.RIGHT)
    text(slide, 0.5, 7.1, 8, 0.3, "PeerDesk  ·  LAN/VPN remote desktop for small teams",
         size=10, color=MUTED)


def actor(slide, cx, top, label, sub=None, color=INK, scale=1.0):
    """UML stick figure; returns (cx, mid_y) anchor for association lines."""
    s = scale
    head = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - 0.17 * s), Inches(top),
                                  Inches(0.34 * s), Inches(0.34 * s))
    head.fill.solid()
    head.fill.fore_color.rgb = WHITE
    head.line.color.rgb = color
    head.line.width = Pt(2)
    head.shadow.inherit = False
    neck, hip = top + 0.34 * s, top + 0.85 * s
    line(slide, cx, neck, cx, hip, color=color, width=2)
    line(slide, cx - 0.3 * s, neck + 0.15 * s, cx + 0.3 * s, neck + 0.15 * s, color=color, width=2)
    line(slide, cx, hip, cx - 0.25 * s, hip + 0.4 * s, color=color, width=2)
    line(slide, cx, hip, cx + 0.25 * s, hip + 0.4 * s, color=color, width=2)
    paras = [[(label, {"bold": True, "color": color})]]
    if sub:
        paras.append([(sub, {"size": 10, "color": MUTED})])
    text(slide, cx - 0.95, hip + 0.42 * s, 1.9, 0.6, paras, size=12, align=PP_ALIGN.CENTER)
    return cx, top + 0.6 * s


def usecase(slide, x, y, w, h, title, tag, fill=WHITE, color=INK, dash=False):
    box(slide, x, y, w, h, fill=fill, line=color, shape=MSO_SHAPE.OVAL, dash=dash, line_w=1.5,
        runs=[[(title, {"bold": True, "color": color})],
              [(tag, {"size": 9.5, "color": MUTED})]], size=12.5)
    return {"l": (x, y + h / 2), "r": (x + w, y + h / 2)}


def bullets(slide, x, y, w, h, items, size=13, color=INK, gap=4):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        segs = item if isinstance(item, list) else [(item, {})]
        r = p.add_run()
        r.text = "▪  "
        r.font.size = Pt(size)
        r.font.color.rgb = segs[0][1].get("bullet", TEAL)
        r.font.name = FONT
        for seg, opts in segs:
            r = p.add_run()
            r.text = seg
            r.font.name = opts.get("font", FONT)
            r.font.size = Pt(opts.get("size", size))
            r.font.bold = opts.get("bold", False)
            r.font.color.rgb = opts.get("color", color)
    return tb


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body


# =====================================================================
# Slide 1 — Overview
# =====================================================================
s = prs.slides.add_slide(BLANK)
header(s, 1, "PeerDesk — Self-hosted Remote Desktop",
       "View and control an Ubuntu workstation from macOS or Ubuntu, over LAN/VPN — no cloud, no licence")

text(s, 0.5, 1.3, 7.6, 0.4, "The problem", size=18, bold=True, color=NAVY)
bullets(s, 0.5, 1.75, 7.6, 1.6, [
    "Small team pays for, or fights with, commercial remote-desktop tools for simple LAN/VPN control",
    "Sessions drop and someone has to walk over and restart the host agent",
    "Naive screen capture feels laggy; Wayland/permission failures are silent",
])

text(s, 0.5, 3.25, 7.6, 0.4, "What PeerDesk delivers (v1)", size=18, bold=True, color=NAVY)
bullets(s, 0.5, 3.7, 7.6, 2.0, [
    [("Two native C++ binaries: ", {"bold": True}), ("peerdesk-server", {"font": "Consolas", "size": 12}),
     (" on the Ubuntu host, ", {}), ("peerdesk-client", {"font": "Consolas", "size": 12}),
     (" (Qt6) on the viewer", {})],
    [("Secure by default: ", {"bold": True}),
     ("TLS 1.2+ on every byte; Argon2id + HMAC challenge-response — the password never crosses the wire", {})],
    [("Live view + full control: ", {"bold": True}),
     ("primary X11 display streamed; mouse & keyboard injected with coordinate mapping", {})],
    [("Resilient: ", {"bold": True}),
     ("one viewer per host; disconnect/reconnect without restarting the server", {})],
])

# success metric callout
box(s, 0.5, 5.8, 7.6, 1.1, fill=TEAL_BG, line=TEAL, line_w=1.25,
    runs=[[("Success metric  ", {"bold": True, "color": TEAL}),
           ("A teammate types IP, port, username, password → sees the remote Ubuntu screen → "
            "moves the mouse and types → disconnects and connects again ", {}),
           ("without restarting the server.", {"bold": True})]],
    size=13, align=PP_ALIGN.LEFT)

# scope panel
box(s, 8.6, 1.3, 4.25, 5.6, fill=LIGHT, line=None)
text(s, 8.85, 1.42, 3.8, 0.4, "Scope", size=18, bold=True, color=NAVY)
text(s, 8.85, 1.85, 3.8, 0.3, "IN v1", size=11, bold=True, color=TEAL)
bullets(s, 8.85, 2.12, 3.8, 1.9, [
    "TLS session, Argon2id + HMAC auth",
    "Primary display (X11 or --synthetic)",
    "Mouse + keyboard, scaled coordinates",
    "One viewer; reconnect without restart",
], size=12, gap=2)
text(s, 8.85, 3.9, 3.8, 0.3, "OUT OF SCOPE v1", size=11, bold=True, color=RED)
bullets(s, 8.85, 4.17, 3.8, 2.6, [
    [("File transfer, audio, clipboard", {"bullet": RED})],
    [("Wayland host, multi-monitor", {"bullet": RED})],
    [("NAT relay / “connect from anywhere”", {"bullet": RED})],
    [("Mac/Windows as host; mobile/web", {"bullet": RED})],
    [("Saved profiles, view-only role", {"bullet": RED})],
], size=12, gap=2)

notes(s, "PeerDesk (repo nickname 'TeamViewer') is a hobby/internal tool. Deploy model: self-hosted "
         "server on the Ubuntu host; clients reach IP:port directly over LAN or VPN. Default port 4473.")

# =====================================================================
# Slide 2 — Actors & Use Cases
# =====================================================================
s = prs.slides.add_slide(BLANK)
header(s, 2, "Actors & Use Cases",
       "Three personas, two roles — once authenticated, a viewer has full control")

# system boundary
box(s, 3.2, 1.3, 6.85, 5.25, fill=None, line=NAVY, shape=MSO_SHAPE.RECTANGLE, line_w=1.75)
text(s, 3.3, 1.33, 4, 0.35, "PeerDesk system", size=13, bold=True, color=NAVY)

LX, RX, UW, UH = 3.45, 6.8, 3.05, 0.8
uc3 = usecase(s, LX, 1.8, UW, UH, "Connect & authenticate", "J1 · TLS + challenge", color=BLUE)
uc4 = usecase(s, LX, 2.78, UW, UH, "View remote screen", "J2 · live primary display", color=BLUE)
uc5 = usecase(s, LX, 3.8, UW, UH, "Control mouse & keyboard", "J2 · mapped to host pixels", color=BLUE)
uc6 = usecase(s, LX, 4.8, UW, UH, "Disconnect & reconnect", "J3 · re-auth, no server restart", color=BLUE)
uc1 = usecase(s, RX, 1.6, UW, UH, "Start host server", "J0 · bind port, X11 check", color=ORANGE)
uc2 = usecase(s, RX, 2.5, UW, UH, "Issue logins", "J0 · Argon2id hashes on disk", color=ORANGE)

# deferred use cases
text(s, 3.5, 5.72, 1.5, 0.3, "Deferred (D1):", size=10.5, bold=True, color=MUTED)
for i, (t, tag) in enumerate([("Saved profiles", "J4"), ("Wayland host", "J5"), ("Multi-monitor", "J6")]):
    box(s, 4.75 + i * 1.75, 5.68, 1.65, 0.5, fill=None, line=LINE, shape=MSO_SHAPE.OVAL, dash=True,
        runs=[[(f"{tag} {t}", {"bold": True, "color": MUTED})]], size=10)

# Viewer (abstract) + Sam / Riley generalisation
vx, vy = actor(s, 2.25, 2.95, "Viewer", "any authenticated user", color=BLUE)
for u in (uc3, uc4, uc5, uc6):
    line(s, vx + 0.3, vy, u["l"][0], u["l"][1], color=BLUE, width=1.25)
sx, sy = actor(s, 0.75, 1.45, "Sam", "Viewer · macOS", color=BLUE, scale=0.75)
rx, ry = actor(s, 0.75, 4.35, "Riley", "Viewer · Ubuntu", color=BLUE, scale=0.75)
line(s, sx + 0.25, sy + 0.15, vx - 0.3, vy - 0.15, color=BLUE, width=1.25, tail="triangle")
line(s, rx + 0.25, ry - 0.05, vx - 0.3, vy + 0.25, color=BLUE, width=1.25, tail="triangle")

# Host operator
jx, jy = actor(s, 11.1, 1.45, "Jordan", "Host operator", color=ORANGE)
for u in (uc1, uc2):
    line(s, jx - 0.3, jy, u["r"][0], u["r"][1], color=ORANGE, width=1.25)

# Secondary system actor: X11 desktop
box(s, 10.45, 4.25, 1.9, 0.95, fill=ORANGE_BG, line=ORANGE, shape=MSO_SHAPE.RECTANGLE,
    runs=[[("«system»", {"size": 9.5, "color": MUTED})],
          [("Ubuntu X11 desktop", {"bold": True, "color": ORANGE})],
          [("capture · XTest inject", {"size": 9.5, "color": MUTED})]], size=12)
for u in (uc4, uc5):
    line(s, u["r"][0], u["r"][1], 10.45, 4.72, color=ORANGE, width=1.25, dash=True)

# footer rules
box(s, 0.35, 6.68, 12.6, 0.38, fill=LIGHT, line=None,
    runs=[[("Rules:  ", {"bold": True, "color": NAVY}),
           ("one active viewer per host — a second is told “Host already has a viewer”   ·   "
            "every reconnect re-authenticates   ·   one person can wear both hats", {})]],
    size=11, align=PP_ALIGN.LEFT)

notes(s, "Sam and Riley are specialisations of the Viewer actor; their capabilities are identical. "
         "Jordan owns J0 (start server, issue logins) and benefits from J3 — the server keeps listening "
         "after a viewer leaves. The Ubuntu X11 desktop is a secondary system actor: the server captures "
         "it and injects input into it via XTest.")

# =====================================================================
# Slide 3 — System architecture
# =====================================================================
s = prs.slides.add_slide(BLANK)
header(s, 3, "System Architecture",
       "Two native C++20 binaries sharing one wire-protocol library — a peer session, not a REST API")

# Viewer zone
box(s, 0.35, 1.3, 4.35, 4.35, fill=BLUE_BG, line=None)
text(s, 0.5, 1.36, 4.1, 0.35, [[("VIEWER  ", {"bold": True, "color": BLUE}),
                                ("peerdesk-client · macOS / Ubuntu", {"size": 11, "color": MUTED})]], size=13)
cb = dict(fill=WHITE, line=BLUE, line_w=1.25, align=PP_ALIGN.LEFT, size=11)
box(s, 0.55, 1.8, 3.95, 0.8, runs=[[("MainWindow", {"bold": True, "color": BLUE}), ("  Qt6 Widgets", {"color": MUTED})],
                                    "Connection form (IP, port, user, password), status & errors"], **cb)
box(s, 0.55, 2.8, 3.95, 0.8, runs=[[("VideoSurface", {"bold": True, "color": BLUE})],
                                    "Paints frames · captures mouse/keys · maps widget → host coords"], **cb)
box(s, 0.55, 3.8, 1.75, 0.7, runs=[[("keymap", {"bold": True, "color": BLUE})], "Qt key → X11 keysym"], **cb)
box(s, 0.55, 4.7, 3.95, 0.8, runs=[[("SessionWorker", {"bold": True, "color": BLUE}), ("  own QThread", {"color": MUTED})],
                                    "TLS conn · auth · JPEG decode · ping · input queue"], **cb)
arrow(s, 2.52, 2.6, 2.52, 2.8, color=BLUE, both=True)
arrow(s, 1.42, 3.6, 1.42, 3.8, color=BLUE)
arrow(s, 3.4, 3.6, 3.4, 4.7, color=BLUE, both=True)
text(s, 3.45, 3.95, 1.2, 0.5, "signals /\nslots", size=9.5, color=MUTED)

# Host zone
box(s, 8.65, 1.3, 4.35, 4.35, fill=ORANGE_BG, line=None)
text(s, 8.8, 1.36, 4.1, 0.35, [[("HOST  ", {"bold": True, "color": ORANGE}),
                                ("peerdesk-server · Ubuntu (X11)", {"size": 11, "color": MUTED})]], size=13)
hb = dict(fill=WHITE, line=ORANGE, line_w=1.25, align=PP_ALIGN.LEFT, size=11)
box(s, 8.85, 1.8, 3.95, 0.8, runs=[[("HostServer", {"bold": True, "color": ORANGE}), ("  accept loop", {"color": MUTED})],
                                    "TLS listener · handshake · atomic one-session gate"], **hb)
box(s, 8.85, 2.8, 1.9, 0.9, runs=[[("UserStore", {"bold": True, "color": ORANGE})], "Argon2id hashes + self-signed cert"], **hb)
box(s, 10.9, 2.8, 1.9, 0.9, runs=[[("Session thread", {"bold": True, "color": ORANGE})], "input loop · Ping/Pong · 8 s timeout"], **hb)
box(s, 8.85, 3.9, 1.9, 1.6, runs=[[("Capture thread", {"bold": True, "color": ORANGE})],
                                   "X11Capture or SyntheticCapture → JPEG @ ~10 fps"], **hb)
box(s, 10.9, 3.9, 1.9, 1.6, runs=[[("InputInject", {"bold": True, "color": ORANGE})],
                                   "XTest: mouse move / button / wheel, keys"], **hb)
arrow(s, 9.8, 2.6, 9.8, 2.8, color=ORANGE)
arrow(s, 11.85, 2.6, 11.85, 2.8, color=ORANGE)
arrow(s, 11.85, 3.7, 11.85, 3.9, color=ORANGE)
arrow(s, 10.9, 3.5, 9.8, 3.9, color=ORANGE)

# Network channel
box(s, 4.95, 1.3, 3.45, 4.35, fill=TEAL_BG, line=None)
text(s, 5.05, 1.36, 3.3, 0.6, [[("TLS 1.2+ over TCP", {"bold": True, "color": TEAL})],
                               [("port 4473 · length-prefixed frames", {"size": 10.5, "color": MUTED})]],
     size=13, align=PP_ALIGN.CENTER)
arrow(s, 4.6, 2.45, 8.75, 2.45, color=BLUE, width=2.25)
text(s, 5.05, 2.0, 3.3, 0.45, "Hello · AuthResponse", size=11, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
arrow(s, 8.75, 3.35, 4.6, 3.35, color=ORANGE, width=2.25)
text(s, 5.05, 2.9, 3.3, 0.45, "AuthChallenge · AuthOk / AuthFail", size=11, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)
arrow(s, 8.75, 4.25, 4.6, 4.25, color=ORANGE, width=3.5)
text(s, 5.05, 3.8, 3.3, 0.45, "VideoFrame (JPEG) · Pong", size=11, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)
arrow(s, 4.6, 5.1, 8.75, 5.1, color=BLUE, width=2.25)
text(s, 5.05, 4.65, 3.3, 0.45, "Mouse · Key · Ping · Disconnect", size=11, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
text(s, 5.05, 5.18, 3.3, 0.4, "input has its own path — never queued behind video",
     size=9.5, color=MUTED, align=PP_ALIGN.CENTER)

# Shared library band
box(s, 0.35, 5.85, 12.65, 1.15, fill=PURPLE_BG, line=None)
text(s, 0.5, 5.9, 8, 0.35, [[("SHARED  ", {"bold": True, "color": PURPLE}),
                             ("libpeerdesk_shared · linked into client, server and tests", {"size": 11, "color": MUTED})]], size=13)
mods = [("protocol", "12 message types, pack/unpack"), ("tls", "OpenSSL listener / conn"),
        ("auth", "Argon2id · HMAC-SHA256"), ("jpeg", "libjpeg encode/decode"),
        ("map", "widget ↔ host coords"), ("bytes", "endian helpers")]
for i, (m, d) in enumerate(mods):
    box(s, 0.55 + i * 2.07, 6.28, 1.95, 0.62, fill=WHITE, line=PURPLE, line_w=1.25, size=10.5,
        runs=[[(m, {"bold": True, "color": PURPLE, "font": "Consolas"})], [(d, {"size": 9.5, "color": MUTED})]])

notes(s, "Ownership (DECISIONS B1/B2/B4): the server owns credentials, session occupancy, capture, encode and "
         "inject. The client owns the connection UI, decode, render, local input capture and coordinate mapping. "
         "shared/ owns wire types, TLS helpers, Argon2/HMAC and coordinate mapping. There is no database — "
         "just a users file of salted Argon2id hashes plus a self-signed cert in the server data dir.")

# =====================================================================
# Slide 4 — Session lifecycle & security
# =====================================================================
s = prs.slides.add_slide(BLANK)
header(s, 4, "Session Lifecycle & Security",
       "Journeys J1 → J2 → J3 on the wire: authenticate, stream and control, then release the session")

lanes = [("Viewer client", 1.45, BLUE), ("Host server", 5.2, ORANGE), ("X11 desktop", 7.75, ORANGE)]
for name, cx, col in lanes:
    box(s, cx - 0.9, 1.25, 1.8, 0.45, fill=col, line=None, runs=[[(name, {"bold": True, "color": WHITE})]], size=12)
    line(s, cx, 1.7, cx, 6.95, color=LINE, width=1.25, dash=True)

C, S, X = 1.45, 5.2, 7.75
rows = [
    (C, S, "TLS handshake (self-signed cert)", TEAL, False),
    (C, S, "Hello {version, username}", BLUE, False),
    (S, C, "AuthChallenge {salt, nonce, Argon2 cost}", ORANGE, False),
    (C, S, "AuthResponse {HMAC(Argon2id hash, nonce)}", BLUE, False),
    (S, C, "AuthOk {w,h}  |  AuthFail {BadCreds | Busy}", ORANGE, False),
    (S, X, "grab frame", ORANGE, False),
    (S, C, "VideoFrame {w, h, JPEG}   ↻ ~10 fps", ORANGE, False),
    (C, S, "Mouse {action, x, y} · Key {keysym}", BLUE, False),
    (S, X, "XTest inject", ORANGE, False),
    (C, S, "Ping ⇄ Pong  (drop after 8 s silence)", TEAL, True),
    (C, S, "Disconnect → release session, keep listening", BLUE, False),
]
y = 2.0
for a, b, label, col, both in rows:
    arrow(s, a, y, b, y, color=col, width=1.5, both=both)
    lx, lw = min(a, b) + 0.05, abs(b - a) - 0.1
    text(s, lx - 0.1, y - 0.3, lw + 0.2, 0.28, label, size=10, color=col, bold=True, align=PP_ALIGN.CENTER)
    y += 0.44

# phase brackets
for label, y0, y1 in [("J1", 1.8, 4.1), ("J2", 4.1, 5.95), ("J3", 5.95, 6.9)]:
    box(s, 0.1, y0 + 0.04, 0.5, y1 - y0 - 0.08, fill=LIGHT, line=None, runs=[[(label, {"bold": True, "color": NAVY})]], size=10)

# Right sidebar
box(s, 8.75, 1.25, 4.25, 5.75, fill=LIGHT, line=None)
text(s, 8.95, 1.35, 3.9, 0.4, "Security design", size=16, bold=True, color=TEAL)
bullets(s, 8.95, 1.78, 3.9, 2.6, [
    "Password never on the wire — only an HMAC proof over a fresh 32-byte nonce",
    "Host stores salted Argon2id hashes only (no plaintext)",
    "Constant-time compare (CRYPTO_memcmp)",
    "Unknown users get a deterministic fake salt → no username enumeration",
    "Every reconnect re-authenticates; no resume tokens",
], size=11.5, gap=3)
text(s, 8.95, 4.35, 3.9, 0.4, "Concurrency", size=16, bold=True, color=TEAL)
bullets(s, 8.95, 4.78, 3.9, 1.3, [
    "Server: accept loop → session thread → capture thread; send mutex shared",
    "Client: network in SessionWorker thread, UI thread only paints",
    "Atomic session gate → second viewer is rejected, not stolen",
], size=11.5, gap=3)
box(s, 8.95, 6.2, 3.85, 0.65, fill=WHITE, line=RED, line_w=1,
    runs=[[("Demo gap: ", {"bold": True, "color": RED}),
           ("client does not verify the server cert (SSL_VERIFY_NONE) — pin/verify before production", {})]],
    size=10.5, align=PP_ALIGN.LEFT)

notes(s, "Auth is challenge-response: server sends salt + nonce + Argon2 cost params; the client derives "
         "Argon2id(password, salt) locally and returns HMAC-SHA256(hash, nonce). The server compares against "
         "its stored hash in constant time. Failure reasons are specific (bad credentials vs busy vs protocol) "
         "so the viewer knows whether to check VPN, password, or wait.")

# =====================================================================
# Slide 5 — Tech stack, status & roadmap
# =====================================================================
s = prs.slides.add_slide(BLANK)
header(s, 5, "Tech Stack, Status & Roadmap",
       "A working LAN demo today; production stack locked 2026-08-30")

text(s, 0.5, 1.25, 7.8, 0.4, "Stack: demo build vs locked production target", size=16, bold=True, color=NAVY)
rows = [
    ("Layer", "Demo (in repo, runs today)", "Production (locked)"),
    ("Client UI", "Qt6 Widgets", "Qt6 Widgets  ✓ same"),
    ("Transport", "TLS 1.2+ TCP (OpenSSL)", "TLS 1.2+ TCP  ✓ same"),
    ("Wire format", "Packed structs", "Protobuf schemas"),
    ("Video codec", "JPEG (libjpeg)", "H.264 — FFmpeg x264 / VAAPI"),
    ("Capture", "X11 XGetImage / synthetic", "X11 XShm + XDamage"),
    ("Input inject", "XTest", "XTest  ✓ same"),
    ("Auth", "Argon2id + HMAC-SHA256", "Argon2id + HMAC  ✓ same"),
    ("Build & ship", "CMake · run_demo.sh", "CMake + vcpkg/Conan · .app / .deb / AppImage"),
]
tbl = s.shapes.add_table(len(rows), 3, Inches(0.5), Inches(1.7), Inches(7.8), Inches(4.3)).table
for c, w in enumerate((1.7, 2.8, 3.3)):
    tbl.columns[c].width = Inches(w)
for r, row in enumerate(rows):
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.text = ""
        cell.margin_left = cell.margin_right = Inches(0.08)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = val
        run.font.name = FONT
        run.font.size = Pt(12 if r else 12.5)
        run.font.bold = r == 0 or c == 0
        changed = r > 0 and c == 2 and "same" not in val
        run.font.color.rgb = WHITE if r == 0 else (TEAL if changed else INK)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY if r == 0 else (LIGHT if r % 2 else WHITE)
text(s, 0.5, 6.05, 7.8, 0.3, [[("Teal", {"bold": True, "color": TEAL}),
                               (" = changes for production. The JPEG / packed-struct demo is throwaway, not the starter.", {})]],
     size=10.5, color=MUTED)

# roadmap
text(s, 8.75, 1.25, 4.2, 0.4, "Roadmap", size=16, bold=True, color=NAVY)
phases = [
    ("D0 · done (demo)", TEAL, TEAL_BG, "J0 Server ready · J1 Connect & auth\nJ2 View & control · J3 Reconnect"),
    ("Next · production", BLUE, BLUE_BG, "Re-build J0–J3 on H.264 + Protobuf,\nvcpkg/Conan deps, desktop installers"),
    ("D1 · deferred", ORANGE, ORANGE_BG, "J4 Saved profiles · J5 Wayland\n(PipeWire / portal) · J6 Multi-monitor"),
    ("Out of scope", RED, RGBColor(0xFB, 0xEB, 0xEB), "J7 NAT relay · files · audio\nclipboard · Mac host · mobile/web"),
]
y = 1.72
for i, (title, col, bg, body) in enumerate(phases):
    box(s, 8.75, y, 4.2, 0.92, fill=bg, line=col, line_w=1.25, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        runs=[[(title, {"bold": True, "color": col, "size": 12.5})]] + [[(ln, {"size": 11})] for ln in body.split("\n")])
    if i < len(phases) - 1:
        arrow(s, 10.85, y + 0.92, 10.85, y + 1.05, color=LINE)
    y += 1.05

# try it
box(s, 0.5, 6.4, 12.45, 0.62, fill=NAVY, line=None, align=PP_ALIGN.LEFT,
    runs=[[("Try it   ", {"bold": True, "color": WHITE}),
           ("cmake -S . -B build && cmake --build build -j  ·  ./build/peerdesk-smoke  ·  ./scripts/run_demo.sh",
            {"font": "Consolas", "size": 10.5, "color": RGBColor(0xC9, 0xD3, 0xE3)}),
           ("    login jordan / peerdesk", {"size": 11, "color": WHITE})]],
    size=12)

notes(s, "peerdesk-smoke covers protocol round-trips, auth, the busy-session rejection and reconnect. "
         "run_demo.sh starts the server in synthetic mode (canvas, no inject) by default, or real X11 with "
         "'x11'. Wayland hosts fail fast with 'X11 required for v1'.")

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(f"wrote {OUT}")
